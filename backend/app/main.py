import io
import json
import os
import re
import secrets
import time
import uuid
import zipfile
from pathlib import PurePosixPath
from typing import Optional
from urllib.parse import urlencode

import httpx
from dotenv import load_dotenv
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import RedirectResponse, StreamingResponse, JSONResponse
from pydantic import BaseModel, Field
from starlette.middleware.sessions import SessionMiddleware

load_dotenv()
app = FastAPI(title="AI Project Builder API", version="1.0.0")
secret = os.getenv("SECRET_KEY", "")
if len(secret) < 32:
    # Dev convenience only; set SECRET_KEY explicitly for persistent sessions.
    secret = secret or secrets.token_urlsafe(48)
app.add_middleware(
    SessionMiddleware, secret_key=secret, same_site="lax",
    https_only=os.getenv("COOKIE_SECURE", "false").lower() == "true",
    max_age=60 * 60 * 8,
)
app.add_middleware(
    CORSMiddleware,
    allow_origins=[os.getenv("FRONTEND_URL", "http://localhost:5173")],
    allow_credentials=True, allow_methods=["GET", "POST"], allow_headers=["Content-Type"],
)

PROJECTS = {}
# OAuth tokens live server-side; signed session cookies are not encrypted.
GITHUB_TOKENS = {}
MAX_UPLOAD = 10 * 1024 * 1024
MAX_FILES = 100
MAX_FILE_BYTES = 1024 * 1024
BLOCKED_PARTS = {".git", "node_modules", "__pycache__", ".venv", "venv", "dist", "build", ".next"}
BLOCKED_NAMES = {".env", ".env.local", ".env.production", "id_rsa", "id_ed25519"}
SECRET_RE = re.compile(r"(api[_-]?key|secret|password|token)\s*[:=]\s*['\"]?([A-Za-z0-9_\-]{20,})", re.I)


class AnalyzeBody(BaseModel):
    topic: str = Field(min_length=5, max_length=20000)
    framework: str = Field(default="React + FastAPI", max_length=100)


class GenerateBody(AnalyzeBody):
    pass


class PublishBody(BaseModel):
    project_id: str
    repo_name: str = Field(min_length=1, max_length=100)
    description: str = Field(default="", max_length=350)
    visibility: str = Field(default="private")
    branch: str = Field(default="main", min_length=1, max_length=100)
    create_new: bool = True
    confirmed: bool = False


def owner_id(request: Request) -> str:
    # Project access is scoped to a signed browser session.
    sid = request.session.get("sid")
    if not sid:
        sid = secrets.token_urlsafe(24)
        request.session["sid"] = sid
    return sid


def safe_path(path: str) -> str:
    if not path or len(path) > 240 or "\\" in path or path.startswith("/"):
        raise HTTPException(400, f"Invalid file path: {path!r}")
    p = PurePosixPath(path)
    if any(part in ("", ".", "..") for part in p.parts):
        raise HTTPException(400, f"Unsafe file path: {path!r}")
    if any(part in BLOCKED_PARTS for part in p.parts) or p.name.lower() in BLOCKED_NAMES:
        raise HTTPException(400, f"Excluded file path: {path!r}")
    if p.suffix.lower() in {".pem", ".key", ".p12", ".pfx"}:
        raise HTTPException(400, "Credential files are not allowed")
    return str(p)


def clean_files(raw) -> dict:
    if not isinstance(raw, dict) or not raw:
        raise HTTPException(502, "AI provider returned no files")
    if len(raw) > MAX_FILES:
        raise HTTPException(502, "Generated project exceeded the file limit")
    result = {}
    for path, content in raw.items():
        path = safe_path(str(path))
        if not isinstance(content, str):
            continue
        if len(content.encode("utf-8")) > MAX_FILE_BYTES:
            raise HTTPException(502, f"Generated file too large: {path}")
        # Avoid accidentally publishing common secret patterns.
        if path.lower().endswith((".env", ".pem", ".key")) or path.lower() in BLOCKED_NAMES:
            continue
        if path.endswith(".env.example"):
            content = re.sub(r"(?im)^(\s*[A-Z0-9_]*(?:KEY|SECRET|TOKEN|PASSWORD)[A-Z0-9_]*\s*=).+$", r"\1", content)
        result[path] = content
    if not result:
        raise HTTPException(502, "No safe project files remained after validation")
    return result


async def extract_upload(upload: UploadFile) -> str:
    name = (upload.filename or "").lower()
    data = await upload.read(MAX_UPLOAD + 1)
    if len(data) > MAX_UPLOAD:
        raise HTTPException(413, "Upload exceeds 10 MB limit")
    try:
        if name.endswith(".txt") or name.endswith(".md"):
            return data.decode("utf-8-sig")
        if name.endswith(".pdf"):
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            return "\n".join(page.extract_text() or "" for page in reader.pages)[:20000]
        if name.endswith(".docx"):
            from docx import Document
            doc = Document(io.BytesIO(data))
            return "\n".join(p.text for p in doc.paragraphs)[:20000]
        if name.endswith((".png", ".jpg", ".jpeg", ".webp")):
            try:
                from PIL import Image
                import pytesseract
                return pytesseract.image_to_string(Image.open(io.BytesIO(data)))[:20000]
            except Exception as exc:
                raise HTTPException(422, "Image OCR failed. Install Tesseract OCR and ensure it is on PATH.") from exc
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, f"Could not parse uploaded file: {exc}") from exc
    raise HTTPException(415, "Supported formats: PDF, DOCX, TXT, MD, PNG, JPG, JPEG, WEBP")


def starter_files(topic: str, framework: str) -> dict:
    # Deterministic fallback for local testing without an AI API key.
    return {
        "README.md": f"# Generated Project\n\n## Topic\n{topic[:1000]}\n\n## Stack\n{framework}\n\n## Run\nSee the frontend and backend folders for setup instructions.\n\n> Starter template generated without an AI provider key. Configure AI_API_KEY for full AI code generation.\n",
        ".gitignore": ".env\n.env.*\n!.env.example\nnode_modules/\n.venv/\n__pycache__/\ndist/\n",
        ".env.example": "AI_API_KEY=\n",
        "backend/requirements.txt": "fastapi\nuvicorn[standard]\n",
        "backend/main.py": "from fastapi import FastAPI\napp = FastAPI(title='Generated Project API')\n\n@app.get('/health')\ndef health():\n    return {'status': 'ok'}\n",
        "frontend/index.html": "<!doctype html><html><head><meta charset='UTF-8'><meta name='viewport' content='width=device-width, initial-scale=1.0'><title>Generated Project</title></head><body><main><h1>Generated Project</h1><p>Configure AI_API_KEY to generate a complete topic-specific application.</p></main></body></html>\n",
    }


async def generate_with_ai(topic: str, framework: str) -> dict:
    api_key = os.getenv("AI_API_KEY", "").strip()
    if not api_key:
        return starter_files(topic, framework)
    base = os.getenv("AI_BASE_URL", "https://api.openai.com/v1").rstrip("/")
    model = os.getenv("AI_MODEL", "gpt-4o-mini")
    prompt = f"""Generate a complete, coherent software project from this specification.
Preferred stack: {framework}
Specification:
{topic}

Return ONLY a JSON object mapping relative file paths to complete UTF-8 file contents. No markdown fences.
Include README.md, .gitignore, .env.example (blank placeholders only), dependency manifests, tests where practical, and frontend/backend/database files as required.
Keep files self-contained and runnable. Never include real credentials, secrets, binaries, lockfiles, node_modules, build output, or .env files.
Maximum 30 files, each under 100 KB. Use secure defaults and explain required environment variables in README."""
    async with httpx.AsyncClient(timeout=float(os.getenv("AI_TIMEOUT_SECONDS", "120"))) as client:
        response = await client.post(
            f"{base}/chat/completions",
            headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
            json={"model": model, "temperature": 0.2, "messages": [
                {"role": "system", "content": "You are a senior software engineer. Output valid JSON only."},
                {"role": "user", "content": prompt},
            ]},
        )
    if response.status_code >= 400:
        raise HTTPException(502, f"AI provider error ({response.status_code}): {response.text[:500]}")
    try:
        content = response.json()["choices"][0]["message"]["content"]
        content = re.sub(r"^```(?:json)?\s*|\s*```$", "", content.strip(), flags=re.I)
        return clean_files(json.loads(content))
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(502, "AI provider response was not valid JSON file mapping") from exc


def get_project(request: Request, project_id: str) -> dict:
    project = PROJECTS.get(project_id)
    if not project or project["owner"] != owner_id(request):
        raise HTTPException(404, "Project not found in this session")
    return project


@app.get("/api/health")
async def health():
    return {"status": "ok", "ai_configured": bool(os.getenv("AI_API_KEY")), "github_configured": bool(os.getenv("GITHUB_CLIENT_ID") and os.getenv("GITHUB_CLIENT_SECRET"))}


@app.post("/api/projects/analyze")
async def analyze(body: AnalyzeBody):
    topic = body.topic.strip()
    if len(topic) < 5:
        raise HTTPException(422, "Please provide a more detailed project specification")
    return {
        "summary": topic[:700],
        "framework": body.framework,
        "suggested_modules": ["Core application", "Validation and error handling", "Tests", "Documentation"],
        "ready": True,
    }


@app.post("/api/projects/generate")
async def generate(body: GenerateBody, request: Request):
    topic = body.topic.strip()
    if len(topic) < 5:
        raise HTTPException(422, "Please provide a project specification")
    files = await generate_with_ai(topic, body.framework)
    project_id = str(uuid.uuid4())
    PROJECTS[project_id] = {
        "id": project_id, "owner": owner_id(request), "topic": topic,
        "framework": body.framework, "files": files, "created_at": int(time.time()),
        "mode": "ai" if os.getenv("AI_API_KEY") else "starter-template",
    }
    return {"project_id": project_id, "file_count": len(files), "files": [{"path": p, "size": len(c.encode())} for p, c in files.items()], "mode": PROJECTS[project_id]["mode"]}


@app.get("/api/projects/{project_id}/files")
async def project_files(project_id: str, request: Request):
    p = get_project(request, project_id)
    return {"project_id": project_id, "files": [{"path": path, "content": content, "size": len(content.encode())} for path, content in p["files"].items()]}


@app.get("/api/projects/{project_id}/download")
async def download_project(project_id: str, request: Request):
    p = get_project(request, project_id)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for path, content in p["files"].items():
            safe = safe_path(path)
            archive.writestr(f"ai-project/{safe}", content)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/zip", headers={"Content-Disposition": f'attachment; filename="ai-project-{project_id[:8]}.zip"'})

@app.post("/api/projects/extract")
async def extract_project_specification(file: UploadFile = File(...)):
    text = await extract_upload(file)
    if not text.strip():
        raise HTTPException(422, "No readable text was found in this file")
    return {"filename": file.filename, "text": text}


@app.get("/api/github/login")
async def github_login(request: Request):
    client_id = os.getenv("GITHUB_CLIENT_ID")
    callback = os.getenv("GITHUB_CALLBACK_URL")
    if not client_id or not os.getenv("GITHUB_CLIENT_SECRET") or not callback:
        raise HTTPException(503, "Configure GitHub OAuth credentials in backend/.env first")
    state = secrets.token_urlsafe(32)
    request.session["github_oauth_state"] = state
    params = {"client_id": client_id, "redirect_uri": callback, "scope": "repo", "state": state}
    return RedirectResponse("https://github.com/login/oauth/authorize?" + urlencode(params))


@app.get("/api/github/callback")
async def github_callback(request: Request, code: Optional[str] = None, state: Optional[str] = None):
    expected = request.session.pop("github_oauth_state", None)
    if not code or not state or not expected or not secrets.compare_digest(state, expected):
        raise HTTPException(400, "Invalid GitHub OAuth state or missing code")
    async with httpx.AsyncClient(timeout=20) as client:
        token_resp = await client.post("https://github.com/login/oauth/access_token", headers={"Accept": "application/json"}, json={
            "client_id": os.getenv("GITHUB_CLIENT_ID"), "client_secret": os.getenv("GITHUB_CLIENT_SECRET"), "code": code,
        })
    if token_resp.status_code >= 400 or not token_resp.json().get("access_token"):
        raise HTTPException(502, "GitHub OAuth token exchange failed")
    sid = owner_id(request)
    GITHUB_TOKENS[sid] = token_resp.json()["access_token"]
    async with httpx.AsyncClient(timeout=20) as client:
        user_resp = await client.get("https://api.github.com/user", headers={"Authorization": f"Bearer {GITHUB_TOKENS[sid]}", "Accept": "application/vnd.github+json"})
    if user_resp.status_code >= 400:
        raise HTTPException(502, "Could not read GitHub account")
    request.session["github_user"] = user_resp.json().get("login")
    return RedirectResponse(os.getenv("FRONTEND_URL", "http://localhost:5173") + "?github=connected")


def github_headers(request: Request) -> dict:
    token = GITHUB_TOKENS.get(owner_id(request))
    if not token:
        raise HTTPException(401, "Connect GitHub first")
    return {"Authorization": f"Bearer {token}", "Accept": "application/vnd.github+json", "X-GitHub-Api-Version": "2022-11-28"}


@app.get("/api/github/repositories")
async def github_repositories(request: Request):
    headers = github_headers(request)
    async with httpx.AsyncClient(timeout=20) as client:
        response = await client.get("https://api.github.com/user/repos?per_page=100&sort=updated", headers=headers)
    if response.status_code >= 400:
        raise HTTPException(response.status_code, f"GitHub API error: {response.text[:400]}")
    return {"repositories": [{"full_name": r["full_name"], "name": r["name"], "private": r["private"], "default_branch": r["default_branch"], "html_url": r["html_url"]} for r in response.json()]}


@app.post("/api/github/publish")
async def github_publish(body: PublishBody, request: Request):
    if not body.confirmed:
        raise HTTPException(400, "Explicit confirmation is required before publishing")
    if body.visibility not in {"public", "private"}:
        raise HTTPException(422, "Visibility must be public or private")
    project = get_project(request, body.project_id)
    headers = github_headers(request)
    async with httpx.AsyncClient(timeout=45) as client:
        if body.create_new:
            create_resp = await client.post("https://api.github.com/user/repos", headers=headers, json={
                "name": body.repo_name, "description": body.description,
                "private": body.visibility == "private", "auto_init": True,
            })
            if create_resp.status_code >= 400:
                detail = create_resp.text[:700]
                raise HTTPException(create_resp.status_code, f"GitHub repository creation failed (check name, permissions, and conflicts): {detail}")
            repo = create_resp.json()
            owner, repo_name = repo["owner"]["login"], repo["name"]
        else:
            parts = body.repo_name.split("/", 1)
            if len(parts) != 2:
                raise HTTPException(422, "Existing repository must be specified as owner/repository")
            owner, repo_name = parts
            check = await client.get(f"https://api.github.com/repos/{owner}/{repo_name}", headers=headers)
            if check.status_code >= 400:
                raise HTTPException(check.status_code, f"Cannot access repository: {check.text[:500]}")
            repo = check.json()
        branch = body.branch.strip()
        if not re.fullmatch(r"[A-Za-z0-9._/-]{1,100}", branch) or ".." in branch:
            raise HTTPException(422, "Invalid target branch name")
        # Find branch ref. For a newly auto-initialized repository, use requested branch if it exists,
        # otherwise the repository's default branch is the safe initial target.
        ref_resp = await client.get(f"https://api.github.com/repos/{owner}/{repo_name}/git/ref/heads/{branch}", headers=headers)
        if ref_resp.status_code >= 400 and body.create_new:
            branch = repo.get("default_branch", "main")
            ref_resp = await client.get(f"https://api.github.com/repos/{owner}/{repo_name}/git/ref/heads/{branch}", headers=headers)
        if ref_resp.status_code >= 400:
            raise HTTPException(ref_resp.status_code, f"Target branch not found: {ref_resp.text[:400]}")
        base_sha = ref_resp.json()["object"]["sha"]
        commit_resp = await client.get(f"https://api.github.com/repos/{owner}/{repo_name}/git/commits/{base_sha}", headers=headers)
        if commit_resp.status_code >= 400:
            raise HTTPException(commit_resp.status_code, "Could not load target branch commit")
        base_tree = commit_resp.json()["tree"]["sha"]
        tree_items = []
        for path, content in project["files"].items():
            safe = safe_path(path)
            blob_resp = await client.post(f"https://api.github.com/repos/{owner}/{repo_name}/git/blobs", headers=headers, json={"content": content, "encoding": "utf-8"})
            if blob_resp.status_code >= 400:
                raise HTTPException(blob_resp.status_code, f"Failed uploading blob {safe}: {blob_resp.text[:300]}")
            tree_items.append({"path": safe, "mode": "100644", "type": "blob", "sha": blob_resp.json()["sha"]})
        tree_resp = await client.post(f"https://api.github.com/repos/{owner}/{repo_name}/git/trees", headers=headers, json={"base_tree": base_tree, "tree": tree_items})
        if tree_resp.status_code >= 400:
            raise HTTPException(tree_resp.status_code, f"Failed creating Git tree: {tree_resp.text[:400]}")
        new_commit = await client.post(f"https://api.github.com/repos/{owner}/{repo_name}/git/commits", headers=headers, json={
            "message": f"Publish generated project: {project['topic'][:60]}",
            "tree": tree_resp.json()["sha"], "parents": [base_sha],
        })
        if new_commit.status_code >= 400:
            raise HTTPException(new_commit.status_code, f"Failed creating commit: {new_commit.text[:400]}")
        update_ref = await client.patch(f"https://api.github.com/repos/{owner}/{repo_name}/git/refs/heads/{branch}", headers=headers, json={"sha": new_commit.json()["sha"], "force": False})
        if update_ref.status_code >= 400:
            raise HTTPException(update_ref.status_code, f"Commit created but branch update failed: {update_ref.text[:500]}")
    return {"success": True, "repository_url": repo.get("html_url", f"https://github.com/{owner}/{repo_name}"), "branch": branch, "commit_sha": new_commit.json()["sha"], "files_published": len(tree_items), "message": "GitHub confirmed the commit and branch update."}


@app.exception_handler(HTTPException)
async def http_exception_handler(request: Request, exc: HTTPException):
    return JSONResponse(status_code=exc.status_code, content={"detail": exc.detail}, headers=exc.headers)
